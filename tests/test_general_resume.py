import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import generate_resume_docx as g
from docx import Document

class GeneralResumeTests(unittest.TestCase):
    def setUp(self):
        self.bank=g.parse_bullet_bank(ROOT/'RESUME_BULLET_BANK.md')
        self.profiles=g.parse_profile_bank(ROOT/'PROFILE_BANK.md')
        self.plan=json.loads((ROOT/'TAILORED_RESUME_PLAN_TEMPLATE.json').read_text())
        self.plan.update(workflow_mode='general-resume',job_search_cycle='2026',date='2026-10-19',company='General',group_cvb_roles_before_bullets=True,cvb_layout_rationale='Spanning technical projects',cvb_functional_description='Research Scientist - Wearable ML')
        self.plan['experience']={'cvb_grouped':['CVB-21',{'id':'CVB-10'},{'id':'CVB-09'}]}

    def test_shared_content_and_public_contact_redaction(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            out,public=g.build_general_resume(self.plan,self.bank,self.profiles,root)
            self.assertEqual(out.parent,root/'job-search/2026/00_General-Resume')
            self.assertEqual(out.name,'2026-10-19_Kevin-Urban_General-Resume.docx')
            md=public.read_text()
            doc=Document(out)
            self.assertIn('Last updated: 2026-10-19',md)
            self.assertNotIn('(973)',md)
            self.assertIn('(973)', '\n'.join(p.text for p in doc.paragraphs))
            self.assertEqual(md.count('Official title:'),3)
            for paragraph in doc.paragraphs[3:]:
                if paragraph.text and paragraph.style.name!='Heading 1':
                    self.assertIn(paragraph.text,md.replace("**", ""))
            self.assertIn('- **SensorNet reusable biosignal modeling package:** Developed SensorNet',md)
            # Publications remain ordinary bullets even when their titles contain colons.
            self.assertIn('- Urban et al. "Rethinking the Polar Cap:',md)
            self.assertNotIn('Evidence /',md)
            self.assertEqual(doc.core_properties.title,'Kevin Urban General Resume')
            self.assertEqual(sorted(p.name for p in out.parent.iterdir()),[out.name])

    def test_public_overwrite_requires_reviewed_baseline(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            public=root/'GENERAL_RESUME.md'
            public.write_text('Manual edits to reconcile.')
            with self.assertRaisesRegex(ValueError,'reconcile'):
                g.build_general_resume(self.plan,self.bank,self.profiles,root)
            self.assertFalse((root/'job-search').exists())
            self.plan['public_resume_base_sha256']=hashlib.sha256(public.read_bytes()).hexdigest()
            g.build_general_resume(self.plan,self.bank,self.profiles,root)
            with self.assertRaisesRegex(ValueError,'reconcile'):
                g.build_general_resume(self.plan,self.bank,self.profiles,root)

    def test_invalid_paths_and_public_contacts_leave_no_artifacts(self):
        for updates in [{'job_search_cycle':'../private'}, {'date':'tomorrow'},
                        {'public_contact_email':'email@example.com | private phone'}]:
            with self.subTest(updates=updates),tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp)
                plan=dict(self.plan,**updates)
                with self.assertRaises(ValueError):
                    g.build_general_resume(plan,self.bank,self.profiles,root)
                self.assertEqual(list(root.iterdir()),[])

if __name__=='__main__':
    unittest.main()
