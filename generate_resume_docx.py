#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - exercised only in missing envs
    raise SystemExit(
        "Missing python-docx. Install dependencies with `python -m pip install -r requirements.txt`."
    ) from exc


REPO_ROOT = Path(__file__).resolve().parent
DEFAULT_BULLET_BANK = REPO_ROOT / "RESUME_BULLET_BANK.md"
DEFAULT_PROFILE_BANK = REPO_ROOT / "PROFILE_BANK.md"

FONT = "Calibri"
INK = RGBColor(31, 31, 31)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(82, 82, 82)


@dataclass(frozen=True)
class BankEntry:
    id: str
    label: str
    text: str
    scores: dict[str, int]
    guidance: str


@dataclass(frozen=True)
class Profile:
    name: str
    headline: str
    text: str


@dataclass(frozen=True)
class Role:
    key: str
    organization: str
    title: str
    location: str
    dates: str
    summary: str | None = None


ROLE_ORDER = [
    Role(
        "podimetrics",
        "Podimetrics",
        "Principal Data Scientist",
        "Remote",
        "Dec 2024 - Apr 2026",
        "Adaptive statistical estimation, computer vision, and physics-informed modeling for FDA-cleared remote-monitoring products using thermal arrays and pressure sensors for diabetic foot ulcer detection and fall-risk assessment.",
    ),
    Role(
        "cvb_director",
        "Cohen Veterans Bioscience",
        "Director of Data Science & Digital Health",
        "New York, NY",
        "Jul 2021 - May 2024",
    ),
    Role(
        "cvb_associate_director",
        "Cohen Veterans Bioscience",
        "Associate Director of Data Science - Sensor Analytics",
        "New York, NY",
        "Jun 2019 - Jul 2021",
    ),
    Role(
        "cvb_senior_ds",
        "Cohen Veterans Bioscience",
        "Senior Data Scientist, Early Signal Team",
        "New York, NY",
        "Jun 2018 - Jun 2019",
    ),
    Role(
        "wwe_senior_ds",
        "WWE",
        "Senior Data Scientist",
        "Stamford, CT",
        "Aug 2017 - Jul 2018",
    ),
    Role(
        "wwe_data_scientist",
        "WWE",
        "Data Scientist",
        "Stamford, CT",
        "Oct 2016 - Aug 2017",
    ),
    Role(
        "cstr_research_assistant",
        "Center for Solar-Terrestrial Research",
        "Research Assistant, New Jersey Institute of Technology",
        "Newark, NJ",
        "Aug 2012 - May 2016",
        "PhD research inferred polar-cap electrodynamics from distributed magnetometer, satellite, and spacecraft observations using novel spectral methods for noisy geophysical time series.",
    ),
]

ROLE_ALLOWED_ENTRY_IDS = {
    "cvb_director": {
        "CVB-01",
        "CVB-02",
        "CVB-03",
        "CVB-04",
        "CVB-05",
        "CVB-06",
        "CVB-17",
    },
    "cvb_associate_director": {
        "CVB-07",
        "CVB-08",
        "CVB-09",
        "CVB-10",
        "CVB-11",
        "CVB-12",
        "CVB-18",
    },
    "cvb_senior_ds": {
        "CVB-09",
        "CVB-13",
        "CVB-14",
        "CVB-15",
        "CVB-16",
    },
}

EDUCATION_LINES = [
    ("PhD, Physics", "New Jersey Institute of Technology / Rutgers University, 2016."),
    ("MS, Applied Physics", "Minor in Applied Math, NJIT / Rutgers University, 2010."),
    ("BS, Applied Physics", "Minor in Applied Math, NJIT / Rutgers University, 2008."),
]


def read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except FileNotFoundError as exc:
        raise SystemExit(f"Missing required file: {path}") from exc


def parse_table_entries(markdown: str) -> dict[str, dict[str, Any]]:
    entries: dict[str, dict[str, Any]] = {}
    id_re = re.compile(r"^(CORE|POD|CVB|WWE|MSG|CSTR|INT|SKILL|PUB)-\d+$")
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if len(cells) < 8 or not id_re.match(cells[0]):
            continue
        try:
            scores = {
                "HWD": int(cells[2]),
                "SAR": int(cells[3]),
                "BSP": int(cells[4]),
                "IOT": int(cells[5]),
                "CLP": int(cells[6]),
            }
        except ValueError:
            scores = {}
        entries[cells[0]] = {
            "label": cells[1],
            "scores": scores,
            "guidance": cells[7],
        }
    return entries


def parse_canonical_text(markdown: str) -> dict[str, tuple[str, str]]:
    if "## Canonical Text" not in markdown:
        raise SystemExit("RESUME_BULLET_BANK.md is missing `## Canonical Text`.")
    section = markdown.split("## Canonical Text", 1)[1]
    pattern = re.compile(
        r"^###\s+((?:CORE|POD|CVB|WWE|MSG|CSTR|INT|SKILL|PUB)-\d+):\s*(.*?)\s*$\n+(.*?)(?=^###\s+(?:CORE|POD|CVB|WWE|MSG|CSTR|INT|SKILL|PUB)-\d+:|\Z)",
        re.M | re.S,
    )
    texts: dict[str, tuple[str, str]] = {}
    for match in pattern.finditer(section):
        entry_id = match.group(1).strip()
        label = match.group(2).strip()
        text = re.sub(r"\n{2,}", "\n\n", match.group(3).strip())
        text = " ".join(line.strip() for line in text.splitlines() if line.strip())
        texts[entry_id] = (label, text)
    return texts


def parse_bullet_bank(path: Path) -> dict[str, BankEntry]:
    markdown = read_text(path)
    table_entries = parse_table_entries(markdown)
    canonical_text = parse_canonical_text(markdown)

    bank: dict[str, BankEntry] = {}
    missing_text = sorted(set(table_entries) - set(canonical_text))
    if missing_text:
        raise SystemExit(
            "These IDs are in the scoring tables but missing canonical text: "
            + ", ".join(missing_text)
        )

    for entry_id, meta in table_entries.items():
        label, text = canonical_text[entry_id]
        bank[entry_id] = BankEntry(
            id=entry_id,
            label=label or meta["label"],
            text=text,
            scores=meta.get("scores", {}),
            guidance=meta.get("guidance", ""),
        )
    return bank


def parse_profile_bank(path: Path) -> dict[str, Profile]:
    markdown = read_text(path)
    chunks = re.split(r"^##\s+", markdown, flags=re.M)
    profiles: dict[str, Profile] = {}
    for chunk in chunks[1:]:
        lines = chunk.splitlines()
        if not lines:
            continue
        name = lines[0].strip()
        if name == "Profile Cluster Update Rules":
            continue
        body = "\n".join(lines[1:])
        headline_match = re.search(r"^\*\*Headline:\*\*\s*(.+)$", body, re.M)
        profile_match = re.search(r"^\*\*Profile:\*\*\s*(.+)$", body, re.M)
        if headline_match and profile_match:
            profiles[name] = Profile(
                name=name,
                headline=headline_match.group(1).strip(),
                text=profile_match.group(1).strip(),
            )
    return profiles


def set_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_border(paragraph, color: str = "D9E2F3") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.45)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3.1)
    normal.paragraph_format.line_spacing = 1.055

    for style_name, size, color, before, after in [
        ("Heading 1", 11.3, BLUE, 8, 3),
        ("Heading 2", 10.3, DARK_BLUE, 5, 1),
        ("Heading 3", 9.7, DARK_BLUE, 4, 1),
    ]:
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = FONT
    bullet._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    bullet.font.size = Pt(9.05)
    bullet.paragraph_format.left_indent = Inches(0.24)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(2.35)
    bullet.paragraph_format.line_spacing = 1.035

    ensure_para_style(styles, "Resume Role", size=9.25, color=MUTED, before=0, after=1.4)
    ensure_para_style(
        styles,
        "Resume Org",
        size=10.4,
        color=DARK_BLUE,
        before=4.5,
        after=0,
        bold=True,
    )


def ensure_para_style(
    styles,
    name: str,
    *,
    size: float,
    color: RGBColor,
    before: float,
    after: float,
    bold: bool = False,
) -> None:
    if name not in styles:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles[name]
    style.base_style = styles["Normal"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True


def add_text(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: float | None = None,
    color: RGBColor | None = None,
):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def add_title(doc: Document, headline: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_text(p, "Kevin Urban, PhD", bold=True, size=18, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_text(p, headline, bold=True, size=9.4, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    add_text(
        p,
        "kevin.ddu@gmail.com | Nutley, NJ | linkedin.com/in/machinesaidgo | github.com/krbnite",
        size=9.2,
        color=MUTED,
    )
    set_paragraph_border(p)


def add_section(doc: Document, title: str) -> None:
    p = doc.add_paragraph(title.upper(), style="Heading 1")
    set_paragraph_border(p, "B7C9E2")


def add_org(doc: Document, org: str) -> None:
    doc.add_paragraph(org, style="Resume Org")


def add_role(doc: Document, role: Role) -> None:
    p = doc.add_paragraph(style="Resume Role")
    add_text(p, role.title, bold=True, size=9.25, color=INK)
    add_text(p, f" | {role.location} | ", size=9.25, color=MUTED)
    add_text(p, role.dates, italic=True, size=9.1, color=MUTED)


def add_body(doc: Document, text: str, *, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3.1)
    p.paragraph_format.line_spacing = 1.055
    add_text(p, text, italic=italic, size=9.25)


def add_labeled_bullet(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_after = Pt(2.35)
    p.paragraph_format.line_spacing = 1.035
    add_text(p, f"{label}: ", bold=True, size=9.05)
    add_text(p, text, size=9.05)


def add_simple_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_after = Pt(2.35)
    p.paragraph_format.line_spacing = 1.035
    add_text(p, text, size=9.05)


def add_skill_line(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.2)
    add_text(p, f"{label}: ", bold=True, size=9.05)
    add_text(p, text, size=9.05)


def resolve_entry(spec: Any, bank: dict[str, BankEntry]) -> BankEntry:
    if isinstance(spec, str):
        try:
            return bank[spec]
        except KeyError as exc:
            raise ValueError(f"Unknown bullet-bank ID: {spec}") from exc
    if not isinstance(spec, dict):
        raise ValueError(f"Bullet spec must be a string ID or object, got: {spec!r}")

    entry_id = spec.get("id")
    if entry_id:
        base = bank.get(entry_id)
        if base is None:
            raise ValueError(f"Unknown bullet-bank ID: {entry_id}")
        return BankEntry(
            id=entry_id,
            label=spec.get("label") or base.label,
            text=spec.get("text") or base.text,
            scores=base.scores,
            guidance=base.guidance,
        )

    label = spec.get("label")
    text = spec.get("text")
    if not label or not text:
        raise ValueError(f"Custom bullet requires `label` and `text`: {spec!r}")
    return BankEntry(id="CUSTOM", label=label, text=text, scores={}, guidance="")


def resolve_entries(specs: list[Any], bank: dict[str, BankEntry]) -> list[BankEntry]:
    return [resolve_entry(spec, bank) for spec in specs]


def resolve_output_path(plan: dict[str, Any], out_arg: str | None) -> Path:
    if out_arg:
        path = Path(out_arg)
    elif plan.get("output_path"):
        path = Path(plan["output_path"])
    else:
        date = plan.get("date", "YYYY-MM-DD")
        company = slugify(plan.get("company", "Company"))
        path = Path(f"{date}_Kevin-Urban_Resume_{company}.docx")
    if not path.is_absolute():
        path = REPO_ROOT / path
    return path


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-")
    return slug or "Company"


def load_plan(path: Path) -> dict[str, Any]:
    try:
        return json.loads(read_text(path))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in {path}: {exc}") from exc


def get_profile(plan: dict[str, Any], profiles: dict[str, Profile]) -> tuple[str, str, str]:
    cluster = plan.get("profile_cluster") or plan.get("cluster") or "Default Sensor ML Research Scientist"
    profile = profiles.get(cluster)
    if profile is None and (not plan.get("headline") or not plan.get("profile")):
        available = ", ".join(sorted(profiles))
        raise ValueError(f"Unknown profile cluster `{cluster}`. Available clusters: {available}")
    headline = plan.get("headline") or (profile.headline if profile else "")
    text = plan.get("profile") or plan.get("profile_text") or (profile.text if profile else "")
    if not headline or not text:
        raise ValueError("Plan must provide headline/profile text or a valid profile_cluster.")
    return cluster, headline, text


def validate_plan(plan: dict[str, Any], bank: dict[str, BankEntry], profiles: dict[str, Profile]) -> None:
    if not plan.get("company"):
        raise ValueError("Plan is missing required field `company`.")
    get_profile(plan, profiles)
    for key in ["core_theme_ids", "project_ids", "internship_ids", "skill_ids", "publication_ids"]:
        for spec in plan.get(key, []):
            resolve_entry(spec, bank)
    experience = plan.get("experience", {})
    if not isinstance(experience, dict):
        raise ValueError("Plan field `experience` must be an object keyed by role key.")
    valid_roles = {role.key for role in ROLE_ORDER}
    unknown_roles = sorted(set(experience) - valid_roles)
    if unknown_roles:
        raise ValueError(
            "Unknown experience role key(s): "
            + ", ".join(unknown_roles)
            + ". Valid keys: "
            + ", ".join(sorted(valid_roles))
        )
    for role_key, specs in experience.items():
        if not isinstance(specs, list):
            raise ValueError(f"`experience.{role_key}` must be a list.")
        for spec in specs:
            entry = resolve_entry(spec, bank)
            allowed_ids = ROLE_ALLOWED_ENTRY_IDS.get(role_key)
            allow_override = isinstance(spec, dict) and bool(spec.get("allow_role_override"))
            if (
                allowed_ids is not None
                and entry.id != "CUSTOM"
                and entry.id not in allowed_ids
                and not allow_override
            ):
                raise ValueError(
                    f"{entry.id} is not in the canonical bullet set for `{role_key}`. "
                    "Move it to the correct CVB role or set `allow_role_override: true` "
                    "with an audit explanation if the bullet intentionally spans titles."
                )


def build_docx(plan: dict[str, Any], bank: dict[str, BankEntry], profiles: dict[str, Profile], output_path: Path) -> None:
    _, headline, profile_text = get_profile(plan, profiles)

    doc = Document()
    configure_styles(doc)
    add_title(doc, headline)

    add_section(doc, "Profile")
    for paragraph in split_paragraphs(profile_text):
        add_body(doc, paragraph)

    core_entries = resolve_entries(plan.get("core_theme_ids", []), bank)
    if core_entries:
        add_section(doc, "Core Research Themes")
        for entry in core_entries:
            add_labeled_bullet(doc, entry.label, entry.text)

    experience = plan.get("experience", {})
    if any(experience.get(role.key) for role in ROLE_ORDER):
        add_section(doc, "Experience")
        current_org = None
        role_summaries = plan.get("role_summaries", {})
        group_wwe_roles = bool(plan.get("group_wwe_roles_before_bullets"))
        for role in ROLE_ORDER:
            specs = experience.get(role.key, [])
            if group_wwe_roles and role.key == "wwe_senior_ds":
                wwe_specs = list(experience.get("wwe_senior_ds", [])) + list(
                    experience.get("wwe_data_scientist", [])
                )
                if not wwe_specs:
                    continue
                if role.organization != current_org:
                    add_org(doc, role.organization)
                    current_org = role.organization
                add_role(doc, role)
                data_scientist_role = next(
                    candidate for candidate in ROLE_ORDER if candidate.key == "wwe_data_scientist"
                )
                add_role(doc, data_scientist_role)
                for entry in resolve_entries(wwe_specs, bank):
                    add_labeled_bullet(doc, entry.label, entry.text)
                continue
            if group_wwe_roles and role.key == "wwe_data_scientist":
                continue
            if not specs:
                continue
            if role.organization != current_org:
                add_org(doc, role.organization)
                current_org = role.organization
            add_role(doc, role)
            summary = role_summaries.get(role.key, role.summary)
            if summary:
                add_body(doc, summary, italic=True)
            for entry in resolve_entries(specs, bank):
                add_labeled_bullet(doc, entry.label, entry.text)

    project_entries = resolve_entries(plan.get("project_ids", []), bank)
    if project_entries:
        add_section(doc, str(plan.get("project_section_title", "Selected Research Projects")))
        for entry in project_entries:
            add_labeled_bullet(doc, entry.label, entry.text)

    internship_entries = resolve_entries(plan.get("internship_ids", []), bank)
    if internship_entries:
        add_section(doc, "Internships")
        for entry in internship_entries:
            add_simple_bullet(doc, entry.text)

    skill_entries = resolve_entries(plan.get("skill_ids", []), bank)
    if skill_entries:
        add_section(doc, "Technical Skills")
        for entry in skill_entries:
            add_skill_line(doc, entry.label, entry.text)

    if plan.get("include_education", True):
        add_section(doc, "Education")
        for label, text in EDUCATION_LINES:
            add_skill_line(doc, label, text)

    publication_entries = resolve_entries(plan.get("publication_ids", []), bank)
    if publication_entries:
        add_section(doc, str(plan.get("publication_section_title", "Selected Publications")))
        for entry in publication_entries:
            add_simple_bullet(doc, entry.text)

    doc.core_properties.title = f"Kevin Urban Resume - {plan.get('company', 'Tailored')}"
    doc.core_properties.author = "Kevin Urban"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def split_paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text.strip()) if part.strip()]


def list_ids(bank: dict[str, BankEntry]) -> None:
    for entry_id in sorted(bank):
        entry = bank[entry_id]
        print(f"{entry.id}\t{entry.label}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a tailored Kevin Urban resume DOCX from a JSON plan.")
    parser.add_argument("plan", nargs="?", help="Path to a tailored resume plan JSON file.")
    parser.add_argument("--out", help="Override output DOCX path.")
    parser.add_argument("--bullet-bank", default=str(DEFAULT_BULLET_BANK), help="Path to RESUME_BULLET_BANK.md.")
    parser.add_argument("--profile-bank", default=str(DEFAULT_PROFILE_BANK), help="Path to PROFILE_BANK.md.")
    parser.add_argument("--validate-only", action="store_true", help="Validate the plan without writing a DOCX.")
    parser.add_argument("--list-ids", action="store_true", help="List available bullet-bank IDs and exit.")
    args = parser.parse_args(argv)

    bank = parse_bullet_bank(Path(args.bullet_bank))
    profiles = parse_profile_bank(Path(args.profile_bank))

    if args.list_ids:
        list_ids(bank)
        return 0

    if not args.plan:
        parser.error("the following arguments are required unless --list-ids is used: plan")

    plan = load_plan(Path(args.plan))
    try:
        validate_plan(plan, bank, profiles)
    except ValueError as exc:
        raise SystemExit(f"Plan validation failed: {exc}") from exc

    output_path = resolve_output_path(plan, args.out)
    if args.validate_only:
        cluster, headline, _ = get_profile(plan, profiles)
        print(f"Plan valid for {plan['company']}.")
        print(f"Profile cluster: {cluster}")
        print(f"Headline: {headline}")
        print(f"Output path: {output_path}")
        return 0

    build_docx(plan, bank, profiles, output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
