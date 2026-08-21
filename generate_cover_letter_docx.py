#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import string
from dataclasses import dataclass
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - exercised only in missing envs
    raise SystemExit(
        "Missing python-docx. Install dependencies with `python -m pip install -r requirements.txt`."
    ) from exc


REPO_ROOT = Path(__file__).resolve().parent
DEFAULT_BANK = REPO_ROOT / "COVER_LETTER_BANK.md"

FONT = "Calibri"


@dataclass(frozen=True)
class Module:
    id: str
    label: str
    text: str
    cluster: str
    role: str
    guidance: str


def read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except FileNotFoundError as exc:
        raise SystemExit(f"Missing required file: {path}") from exc


def parse_table_modules(markdown: str) -> dict[str, dict[str, str]]:
    modules: dict[str, dict[str, str]] = {}
    id_re = re.compile(r"^(OPEN|EVID|BRIDGE|CLOSE)-[A-Z0-9]+-\d+$")
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if len(cells) < 4 or not id_re.match(cells[0]):
            continue
        modules[cells[0]] = {
            "cluster": cells[1],
            "role": cells[2],
            "guidance": cells[3],
        }
    return modules


def parse_canonical_text(markdown: str) -> dict[str, tuple[str, str]]:
    if "## Canonical Text" not in markdown:
        raise SystemExit("COVER_LETTER_BANK.md is missing `## Canonical Text`.")
    section = markdown.split("## Canonical Text", 1)[1]
    pattern = re.compile(
        r"^###\s+((?:OPEN|EVID|BRIDGE|CLOSE)-[A-Z0-9]+-\d+):\s*(.*?)\s*$\n+(.*?)(?=^###\s+(?:OPEN|EVID|BRIDGE|CLOSE)-[A-Z0-9]+-\d+:|\Z)",
        re.M | re.S,
    )
    texts: dict[str, tuple[str, str]] = {}
    for match in pattern.finditer(section):
        module_id = match.group(1).strip()
        label = match.group(2).strip()
        text = re.sub(r"\n{2,}", "\n\n", match.group(3).strip())
        text = " ".join(line.strip() for line in text.splitlines() if line.strip())
        texts[module_id] = (label, text)
    return texts


def parse_bank(path: Path) -> dict[str, Module]:
    markdown = read_text(path)
    table_modules = parse_table_modules(markdown)
    canonical_text = parse_canonical_text(markdown)

    missing_text = sorted(set(table_modules) - set(canonical_text))
    if missing_text:
        raise SystemExit(
            "These IDs are in the module table but missing canonical text: "
            + ", ".join(missing_text)
        )

    bank: dict[str, Module] = {}
    for module_id, meta in table_modules.items():
        label, text = canonical_text[module_id]
        bank[module_id] = Module(
            id=module_id,
            label=label,
            text=text,
            cluster=meta["cluster"],
            role=meta["role"],
            guidance=meta["guidance"],
        )
    return bank


def load_plan(path: Path) -> dict[str, Any]:
    try:
        return json.loads(read_text(path))
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in {path}: {exc}") from exc


def company_possessive(company: str) -> str:
    if company.endswith("s") or company.endswith("x") or company.endswith("z"):
        return f"{company}'"
    return f"{company}'s"


def template_vars(plan: dict[str, Any]) -> dict[str, str]:
    variables = {k: str(v) for k, v in plan.get("variables", {}).items()}
    company = str(plan.get("company", "")).strip()
    role_title = str(plan.get("role_title", "")).strip()
    variables.setdefault("company", company)
    variables.setdefault("role_title", role_title)
    variables.setdefault("company_possessive", company_possessive(company))
    return variables


def render_template(text: str, variables: dict[str, str], module_id: str) -> str:
    needed = {
        field_name
        for _, field_name, _, _ in string.Formatter().parse(text)
        if field_name
    }
    missing = sorted(field for field in needed if field not in variables)
    if missing:
        raise ValueError(f"{module_id} is missing variable(s): {', '.join(missing)}")
    rendered = text.format_map(variables)
    unresolved = sorted(set(re.findall(r"\{[A-Za-z0-9_]+\}", rendered)))
    if unresolved:
        raise ValueError(f"{module_id} has unresolved placeholder(s): {', '.join(unresolved)}")
    return rendered


def resolve_module(spec: Any, bank: dict[str, Module], variables: dict[str, str]) -> Module:
    if isinstance(spec, str):
        try:
            base = bank[spec]
        except KeyError as exc:
            raise ValueError(f"Unknown cover-letter module ID: {spec}") from exc
        return Module(
            id=base.id,
            label=base.label,
            text=render_template(base.text, variables, base.id),
            cluster=base.cluster,
            role=base.role,
            guidance=base.guidance,
        )

    if not isinstance(spec, dict):
        raise ValueError(f"Module spec must be a string ID or object: {spec!r}")

    module_id = spec.get("id")
    if module_id:
        base = bank.get(module_id)
        if base is None:
            raise ValueError(f"Unknown cover-letter module ID: {module_id}")
        text = spec.get("text") or base.text
        return Module(
            id=module_id,
            label=spec.get("label") or base.label,
            text=render_template(text, variables, module_id),
            cluster=base.cluster,
            role=base.role,
            guidance=base.guidance,
        )

    label = spec.get("label") or "Custom paragraph"
    text = spec.get("text")
    if not text:
        raise ValueError(f"Custom paragraph requires `text`: {spec!r}")
    return Module(
        id="CUSTOM",
        label=label,
        text=render_template(text, variables, label),
        cluster="Custom",
        role="Custom",
        guidance="",
    )


def resolve_modules(specs: list[Any], bank: dict[str, Module], variables: dict[str, str]) -> list[Module]:
    return [resolve_module(spec, bank, variables) for spec in specs]


def resolve_output_path(plan: dict[str, Any], out_arg: str | None) -> Path:
    if out_arg:
        path = Path(out_arg)
    elif plan.get("output_path"):
        path = Path(plan["output_path"])
    else:
        date = plan.get("date", "YYYY-MM-DD")
        company = slugify(plan.get("company", "Company"))
        path = Path(f"{date}_Kevin-Urban_Cover-Letter_{company}.docx")
    if not path.is_absolute():
        path = REPO_ROOT / path
    return path


def slugify(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", str(value)).strip("-")
    return slug or "Company"


def validate_plan(plan: dict[str, Any], bank: dict[str, Module]) -> None:
    if not plan.get("company"):
        raise ValueError("Plan is missing required field `company`.")
    if not plan.get("role_title"):
        raise ValueError("Plan is missing required field `role_title`.")
    paragraphs = plan.get("paragraphs")
    if not isinstance(paragraphs, list) or not paragraphs:
        raise ValueError("Plan field `paragraphs` must be a non-empty list.")
    variables = template_vars(plan)
    resolve_modules(paragraphs, bank, variables)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.08

    if "Letter Body" not in styles:
        letter_body = styles.add_style("Letter Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        letter_body = styles["Letter Body"]
    letter_body.base_style = normal
    letter_body.font.name = FONT
    letter_body._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    letter_body._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    letter_body.font.size = Pt(11)
    letter_body.paragraph_format.space_after = Pt(10)
    letter_body.paragraph_format.line_spacing = 1.08


def add_paragraph(doc: Document, text: str, *, after: float | None = None) -> None:
    paragraph = doc.add_paragraph(style="Letter Body")
    paragraph.add_run(text)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)


def build_docx(plan: dict[str, Any], bank: dict[str, Module], output_path: Path) -> None:
    variables = template_vars(plan)
    modules = resolve_modules(plan.get("paragraphs", []), bank, variables)

    doc = Document()
    configure_styles(doc)

    if plan.get("include_date"):
        add_paragraph(doc, str(plan.get("date", "")), after=10)

    for line in plan.get("recipient_block", []):
        add_paragraph(doc, str(line), after=0)
    if plan.get("recipient_block"):
        doc.add_paragraph()

    add_paragraph(doc, plan.get("salutation", "Dear Hiring Team,"), after=10)
    for module in modules:
        add_paragraph(doc, module.text, after=10)

    add_paragraph(doc, plan.get("closing", "Sincerely,"), after=0)
    add_paragraph(doc, plan.get("signature", "Kevin Urban, PhD"), after=0)

    doc.core_properties.title = f"Kevin Urban Cover Letter - {plan.get('company', 'Tailored')}"
    doc.core_properties.author = "Kevin Urban"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate a tailored cover-letter DOCX from a JSON plan."
    )
    parser.add_argument("plan", nargs="?", help="Path to a cover-letter plan JSON file.")
    parser.add_argument("--bank", default=str(DEFAULT_BANK), help="Path to COVER_LETTER_BANK.md.")
    parser.add_argument("--out", help="Override output DOCX path.")
    parser.add_argument("--list-ids", action="store_true", help="List available cover-letter module IDs and exit.")
    parser.add_argument("--validate-only", action="store_true", help="Validate the plan without writing a DOCX.")
    args = parser.parse_args()

    bank = parse_bank(Path(args.bank))
    if args.list_ids:
        for module_id in sorted(bank):
            print(f"{module_id}\t{bank[module_id].label}")
        return

    if not args.plan:
        raise SystemExit("A plan JSON path is required unless --list-ids is used.")

    plan = load_plan(Path(args.plan))
    validate_plan(plan, bank)
    output_path = resolve_output_path(plan, args.out)

    if args.validate_only:
        print(f"Plan valid for {plan['company']}.")
        print(f"Paragraph count: {len(plan.get('paragraphs', []))}")
        print(f"Output path: {output_path}")
        return

    build_docx(plan, bank, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
